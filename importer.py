import csv
import re
import sqlite3
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook

DATE_FORMATS = ["%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y"]
NAME_UNIT_RE = re.compile(r"^(.*?)\s*\(([^()]+)\)\s*$")

EXPECTED_ITEMS_COLUMNS = {"id", "name", "unit"}
EXPECTED_INVENTORY_COLUMNS = {"id", "date", "item_id", "amount", "unit"}


def read_csv(filename):
    with open(
        filename,
        "r",
        newline="",
        encoding="utf-8-sig"
    ) as file:
        reader = csv.reader(file)
        return list(reader)

def read_xlsx(filename):
    workbook = load_workbook(
        filename,
        data_only=True
    )

    sheet = workbook.active

    return [
        [cell.value for cell in row]
        for row in sheet.iter_rows()
    ]

def list_importable_files():
    """
    Candidate files for import, newest first. Used instead of a native
    file picker — ft.FilePicker is currently broken in Flet (unresolved
    upstream bug), so the app scans ~/Downloads and lets the user pick
    from a list.
    """
    downloads_dir = Path.home() / "Downloads"

    if not downloads_dir.exists():
        return []

    files = [
        path for path in downloads_dir.iterdir()
        if path.is_file() and path.suffix.lower() in (".csv", ".db")
    ]

    return sorted(files, key=lambda path: path.stat().st_mtime, reverse=True)

def parse_date(value):
    if value is None:
        return None

    value = str(value).strip()

    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(value, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue

    return None

def parse_amount(value):
    if value is None:
        return None

    value = str(value).strip()

    if not value or value.lower() == "no data":
        return None

    try:
        return float(value)
    except ValueError:
        return None

def split_name_unit(label):
    """
    "Flour (kg)" -> ("Flour", "kg"). Plain labels with no "(unit)"
    suffix return an empty guessed unit for the user to fill in.
    """
    label = str(label).strip()
    match = NAME_UNIT_RE.match(label)

    if match:
        return match.group(1).strip(), match.group(2).strip()

    return label, ""

def detect_orientation(rows):
    """
    Guesses whether dates run down the first column (matching this
    app's own exports) or across the header row, by checking which
    axis parses more successfully as dates.
    """
    if len(rows) < 2 or len(rows[0]) < 2:
        return "dates_as_rows"

    header_dates = sum(1 for cell in rows[0][1:] if parse_date(cell))
    first_col_dates = sum(1 for row in rows[1:] if parse_date(row[0]))

    if header_dates > first_col_dates:
        return "dates_as_columns"

    return "dates_as_rows"

def extract_labels(rows, orientation):
    """The header text for whichever axis holds item names, not dates."""
    if len(rows) < 2:
        return []

    if orientation == "dates_as_rows":
        return [str(cell) for cell in rows[0][1:]]

    return [str(row[0]) for row in rows[1:] if row]

def extract_records(rows, orientation):
    """
    Returns a list of (date, item_label, amount) tuples for every
    parseable, non-blank cell, given a confirmed orientation.
    """
    records = []

    if len(rows) < 2:
        return records

    if orientation == "dates_as_rows":
        labels = [str(cell) for cell in rows[0][1:]]

        for row in rows[1:]:
            if not row:
                continue

            date = parse_date(row[0])

            if not date:
                continue

            for label, cell in zip(labels, row[1:]):
                amount = parse_amount(cell)

                if amount is not None:
                    records.append((date, label, amount))
    else:
        dates = [str(cell) for cell in rows[0][1:]]

        for row in rows[1:]:
            if not row:
                continue

            label = str(row[0])

            for date_raw, cell in zip(dates, row[1:]):
                date = parse_date(date_raw)
                amount = parse_amount(cell)

                if date and amount is not None:
                    records.append((date, label, amount))

    return records

def validate_db_schema(filename):
    """
    Only accepts .db files that match this app's exact schema — no
    attempt to guess or adapt a foreign structure.
    """
    try:
        conn = sqlite3.connect(str(filename))

        items_columns = {
            row[1] for row in conn.execute("PRAGMA table_info(items)")
        }
        inventory_columns = {
            row[1] for row in conn.execute("PRAGMA table_info(inventory)")
        }

        conn.close()

        return (
            items_columns == EXPECTED_ITEMS_COLUMNS
            and inventory_columns == EXPECTED_INVENTORY_COLUMNS
        )
    except sqlite3.Error:
        return False

def read_db(filename):
    """
    Returns (items, inventory) from a foreign .db file already
    confirmed valid by validate_db_schema().
    """
    conn = sqlite3.connect(str(filename))

    items = conn.execute("SELECT id, name, unit FROM items").fetchall()
    inventory = conn.execute(
        "SELECT id, date, item_id, amount, unit FROM inventory"
    ).fetchall()

    conn.close()

    return items, inventory